from pathlib import Path
import json,hashlib
import openpyxl
p=Path('files/演算器/新华人寿 宏泰世家（分红型）.xlsx')
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
s=w['宏泰世家终身寿险（分红型）相关利益演示数据']
print('Workbook SHA256',hashlib.sha256(p.read_bytes()).hexdigest())
for r in s.iter_rows(min_row=3,max_row=6,max_col=17,values_only=True):print('HEADER',r)
d=json.loads(Path('data_hongtai.json').read_text(encoding='utf8'));n=0;errors=[]
for row in s.iter_rows(min_row=5,max_col=17,values_only=True):
 if not isinstance(row[0],str)or row[0]not in d:continue
 r=d[row[0]];n+=1
 for field,index in [('base_sa',4),('cv',10),('div_annual',13),('paidup_net',15),('paidup_cv',16)]:
  if r[field]!=row[index]:errors.append({'key':row[0],'field':field,'json':r[field],'excel':row[index]})
print('Matched raw source rows',n,'mapping mismatches',len(errors));print(json.dumps(errors[:3],ensure_ascii=False))
assert n>30000 and not errors
